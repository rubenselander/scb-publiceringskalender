"""Deterministic workbook, collection and change-report metadata extraction."""

import io
import re
from contextlib import closing
from dataclasses import dataclass
from datetime import date
from urllib.parse import urlsplit
from zipfile import BadZipFile

import pdfplumber
from bs4 import BeautifulSoup, Tag
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pdfminer.pdfdocument import PDFException
from pdfminer.psparser import PSException
from pdfplumber.utils.exceptions import PdfminerException

from scb_extract.core import (
    ExtractionContext,
    FetchRequest,
    FetchSnapshot,
    SourceResult,
    absolute_url,
    url_key,
)
from scb_extract.models.common import IndexEntry, PageIndex
from scb_extract.models.downloads_collections import (
    ChangeNotice,
    ChangesReport,
    EconomyCollection,
    EconomyDiagram,
    HvdCollection,
    HvdGroup,
    HvdLink,
    OfficialProductRow,
    OfficialProductWorkbook,
)

ROOT = "https://www.scb.se"
URLS = {
    "official_products": ROOT + "/sam-forum/hem/officiell-statistik/",
    "hvd": ROOT
    + "/vara-tjanster/oppna-data/vardefulla-datamangder-hvd/vardefulla-datamangder--statistik/",
    "economy": ROOT
    + "/hitta-statistik/statistik-efter-amne/ovrigt/allmant/sveriges-ekonomi/",
    "changes": ROOT
    + "/sam-forum/hem/officiell-statistik/andringar-i-den-officiella-statistiken/",
}
HEADERS = [
    "Statistikansvarig myndighet",
    "Produktkod",
    "Produktnamn",
    "Syfte",
    "Anvandare",
    "Publiceringsstatus",
    "Ämnesområde",
    "Statistikområde",
    "Periodicitet",
]
CODE = re.compile(r"^[A-Z]{2}\d{4}$")
SOURCE_ERRORS = (
    OSError,
    ValueError,
    BadZipFile,
    InvalidFileException,
    PDFException,
    PSException,
    PdfminerException,
)


def _text(node: Tag | None) -> str:
    return node.get_text(" ", strip=True) if node else ""


def _html(snapshot: FetchSnapshot) -> BeautifulSoup:
    # Parse bytes so the response's HTML charset participates in decoding.
    return BeautifulSoup(snapshot.content, "lxml")


def _provenance(snapshot: FetchSnapshot, selector: str):
    soup = _html(snapshot)
    canonical = soup.select_one('link[rel="canonical"]')
    return snapshot.provenance(
        canonical=absolute_url(snapshot.final_url, canonical["href"])
        if canonical and canonical.get("href")
        else snapshot.final_url,
        selector=selector,
    )


def parse_workbook(snapshot: FetchSnapshot) -> OfficialProductWorkbook:
    """Preserve original cells and reject incompatible workbook structures."""
    with closing(
        load_workbook(io.BytesIO(snapshot.content), read_only=True, data_only=True)
    ) as book:
        if len(book.worksheets) != 1:
            raise ValueError("Official product workbook must contain exactly one sheet")
        sheet = book.worksheets[0]
        values = list(sheet.iter_rows(values_only=True))
        if not values or list(values[0]) != HEADERS:
            raise ValueError("Official product workbook headers changed")
        rows = []
        for ordinal, cells in enumerate(values[1:], 2):
            if all(value is None for value in cells):
                raise ValueError(f"Unexpected blank workbook row {ordinal}")
            rows.append(
                OfficialProductRow.model_validate(
                    dict(zip(HEADERS, cells, strict=True))
                )
            )
        codes = [row.product_code for row in rows]
        if not rows or len(set(codes)) != len(codes):
            raise ValueError("Empty workbook or duplicate product codes")
        return OfficialProductWorkbook(
            provenance=snapshot.provenance(
                canonical=snapshot.final_url, selector=f"sheet:{sheet.title}"
            ),
            sheet=sheet.title,
            headers=HEADERS,
            total_source_rows=len(rows),
            partial_sample=False,
            rows=rows,
        )


def parse_hvd(snapshot: FetchSnapshot) -> HvdCollection:
    soup = _html(snapshot)
    tables = [
        table
        for table in soup.select("main table")
        if [_text(cell) for cell in table.select("thead td, thead th")]
        == ["Datamängd", "Tabell/sparad fråga"]
    ]
    if len(tables) != 1:
        raise ValueError("Expected one HVD collection table")
    groups = []
    for row in tables[0].select("tbody tr"):
        cells = row.find_all("td", recursive=False)
        if len(cells) != 2 or not _text(cells[0]):
            raise ValueError("Incomplete HVD table row")
        links = []
        for anchor in cells[1].select("a[href]"):
            url = absolute_url(snapshot.final_url, anchor["href"])
            links.append(
                HvdLink(
                    anchor_text=_text(anchor),
                    url=url,
                    target_kind="pxweb_table"
                    if "/pxweb/" in url.lower()
                    else "saved_query"
                    if "/sq/" in url.lower()
                    else "other",
                    context_text=_text(anchor.find_parent("p") or cells[1]) or None,
                )
            )
        groups.append(
            HvdGroup(dataset_label=_text(cells[0]), row_text=_text(row), links=links)
        )
    if not groups:
        raise ValueError("Empty HVD collection")
    return HvdCollection(
        provenance=_provenance(snapshot, "main table"),
        title=_text(soup.select_one("h1")),
        headers=["Datamängd", "Tabell/sparad fråga"],
        total_source_groups=len(groups),
        partial_sample=False,
        groups=groups,
    )


def parse_economy(snapshot: FetchSnapshot) -> EconomyCollection:
    soup = _html(snapshot)
    tables = [
        table
        for table in soup.select("main table")
        if [_text(cell) for cell in table.select("thead th")]
        == ["Namn", "Typ", "Datum"]
    ]
    if len(tables) != 1:
        raise ValueError("Expected one economy diagram inventory")
    entries = []
    for row in tables[0].select("tbody tr"):
        cells = row.find_all("td", recursive=False)
        anchor = cells[0].select_one("a[href]") if cells else None
        if len(cells) != 3 or not anchor:
            raise ValueError("Incomplete economy diagram row")
        entries.append(
            EconomyDiagram(
                title=_text(anchor),
                url=absolute_url(snapshot.final_url, anchor["href"]),
                listed_type=_text(cells[1]),
                listed_date=date.fromisoformat(_text(cells[2]))
                if _text(cells[2])
                else None,
                detail_verified=False,
                detail_provenance=None,
                subtitle=None,
                comments=None,
                source_label=None,
                updated_at=None,
                excel_urls=[],
                image_urls=[],
                official_statistics_mark_present=None,
            )
        )
    if not entries or len({entry.url for entry in entries}) != len(entries):
        raise ValueError("Empty or duplicate-URL economy inventory")
    return EconomyCollection(
        provenance=_provenance(snapshot, "main table[Namn,Typ,Datum]"),
        title=_text(soup.select_one("h1")),
        section="Tabeller och diagram",
        total_source_entries=len(entries),
        partial_sample=False,
        entries=entries,
    )


def parse_economy_detail(
    snapshot: FetchSnapshot, entry: EconomyDiagram
) -> EconomyDiagram:
    soup = _html(snapshot)
    article = soup.select_one("#pageContent article")
    if not article or not article.select_one("h1"):
        raise ValueError("Missing economy diagram article")
    labels = {}
    for label in article.select("span.text-dark-gray"):
        paragraph = label.find_parent("p")
        value = _text(paragraph)
        labels[_text(label)] = value.removeprefix(_text(label)).strip()
    updated = labels.get("Senast uppdaterad")
    if not updated or not labels.get("Källa"):
        raise ValueError("Missing economy source or update date")
    comments_heading = next(
        (h for h in article.select("h2") if _text(h) == "Kommentarer"), None
    )
    comments = comments_heading.find_next_sibling("div") if comments_heading else None
    return entry.model_copy(
        update={
            "title": _text(article.select_one("h1")),
            "detail_verified": True,
            "detail_provenance": _provenance(snapshot, "#pageContent article"),
            "subtitle": _text(article.select_one("p.ingress")) or None,
            "comments": _text(comments) or None,
            "source_label": labels["Källa"],
            "updated_at": date.fromisoformat(updated),
            "excel_urls": [
                absolute_url(snapshot.final_url, a["href"])
                for a in article.select("a[href]")
                if urlsplit(a["href"]).path.lower().endswith((".xls", ".xlsx"))
            ],
            "image_urls": [
                absolute_url(snapshot.final_url, image["src"])
                for image in article.select(".xhtmlText img[src]")
            ],
            "official_statistics_mark_present": bool(
                soup.select('img[src*="SOS-icons"]')
            ),
        }
    )


def parse_changes_index(snapshot: FetchSnapshot) -> PageIndex:
    soup = _html(snapshot)
    entries = []
    for anchor in soup.select("#floatContentBody a[href]"):
        if urlsplit(anchor["href"]).path.lower().endswith(".pdf"):
            url = absolute_url(snapshot.final_url, anchor["href"])
            entries.append(IndexEntry(key=url_key(url), title=_text(anchor), url=url))
    if not entries or len({entry.key for entry in entries}) != len(entries):
        raise ValueError("Empty or duplicate change-report inventory")
    return PageIndex(
        provenance=_provenance(snapshot, "#floatContentBody a[href$='.pdf']"),
        title=_text(soup.select_one("h1")),
        entries=entries,
        complete=True,
    )


def _interpret_notice(cells: list[str], pages: list[int], ordinal: int) -> ChangeNotice:
    authority, code, name, description = cells
    if not authority or not description or (code and not CODE.fullmatch(code)):
        raise ValueError(f"Unresolved change-report row {ordinal}: {cells!r}")
    lowered = description.lower()
    kinds = []
    for kind, pattern in [
        ("new_product", r"ny produkt"),
        ("name", r"produktnamn|namnändring"),
        (
            "content",
            r"ändring av produkt[.: ]|variabel|redovisningsgrupp|tillägg av tabell",
        ),
        ("periodicity", r"periodicitet"),
        ("inactive", r"produkten (?:bli[r]?|är) inaktiv"),
        ("closed", r"avslutad produkt|produkten avslutas"),
        ("status", r"produktstatus"),
        (
            "merge",
            r"slås ihop|slogs.*ihop|gemensam produktsida|sammanslagning av produkter",
        ),
    ]:
        if re.search(pattern, lowered):
            kinds.append(kind)
    planned = bool(
        re.search(r"avsikt|avser att|föreslås|kommer att|planerar|vill ", lowered)
    )
    effective = bool(
        re.search(
            r"ändringen gäller|(?:gäller |produkt )?fr\.?\s*o\.?\s*m\.?\s*\d{4}-\d{2}-\d{2}"
            r"|produkten är inaktiv|avslutad produkt|är officiell statistik"
            r"|har (?:utökats|lagts till|ändrats)",
            lowered,
        )
    )
    timing = (
        "mixed"
        if planned and effective
        else "planned"
        if planned
        else "source_claimed_effective"
        if effective
        else "unclear"
    )
    # Dates are effective only when tied directly to the change, never publication/reference periods.
    date_match = re.search(
        r"(?:ändringen gäller\s*)?(?:fr\.?\s*o\.?\s*m\.?|från och med)\s*(\d{4}-\d{2}-\d{2})",
        lowered,
    )
    effective_date = date.fromisoformat(date_match[1]) if date_match else None
    return ChangeNotice(
        authority=authority,
        product_code=code or None,
        product_name=name or None,
        description=description,
        pages=pages,
        row_ordinal=ordinal,
        change_kinds=kinds or ["other"],
        timing_claim=timing,
        effective_date=effective_date,
        effective_date_text=date_match[0] if date_match else None,
        related_product_codes=list(
            dict.fromkeys(
                c for c in re.findall(r"\b[A-Z]{2}\d{4}\b", description) if c != code
            )
        ),
        interpretation_evidence=description,
        review_required=timing != "source_claimed_effective" or not code,
    )


def parse_changes(snapshot: FetchSnapshot, *, link_label: str = "") -> ChangesReport:
    """Reconstruct wrapped rows using each page's printed column headings.

    pdfplumber's line-table detection misses alternating rows in these PDFs.
    Word geometry retains every column and continuation across page boundaries.
    Unknown layout or an orphan continuation fails the entire report.
    """
    notices = []
    current = None
    pages = []
    page_texts = {}
    with pdfplumber.open(io.BytesIO(snapshot.content)) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            page_texts[page_number] = page.extract_text() or ""
            words = page.extract_words(x_tolerance=2, y_tolerance=2)
            headings = [
                next((w for w in words if w["text"] == label), None)
                for label in ["Myndighet", "Produktkod", "Produktnamn", "Beskrivning"]
            ]
            if (
                not all(headings)
                or max(w["top"] for w in headings) - min(w["top"] for w in headings) > 3
            ):
                raise ValueError(
                    f"Unresolved report column headings on page {page_number}"
                )
            starts = [w["x0"] - 2 for w in headings]
            bottom = max(w["bottom"] for w in headings)
            lines = []
            for word in sorted(
                (w for w in words if w["top"] > bottom + 1),
                key=lambda w: (w["top"], w["x0"]),
            ):
                if not lines or abs(lines[-1][0] - word["top"]) > 2:
                    lines.append((word["top"], []))
                lines[-1][1].append(word)
            for _, line in lines:
                cells = [[], [], [], []]
                for word in sorted(line, key=lambda w: w["x0"]):
                    column = max(
                        (i for i, start in enumerate(starts) if word["x0"] >= start),
                        default=0,
                    )
                    cells[column].append(word["text"])
                values = [" ".join(cell) for cell in cells]
                if re.fullmatch(r"\d+/\d+", " ".join(values).strip()):
                    continue
                if values[1]:
                    if not CODE.fullmatch(values[1]) or not values[0]:
                        raise ValueError(
                            f"Unresolved code cell on page {page_number}: {values}"
                        )
                    if current is not None:
                        notices.append(
                            _interpret_notice(current, pages, len(notices) + 1)
                        )
                    current, pages = values, [page_number]
                else:
                    if current is None:
                        raise ValueError(
                            f"Orphan row continuation on page {page_number}"
                        )
                    for i, value in enumerate(values):
                        if value:
                            current[i] += " " + value
                    if page_number not in pages:
                        pages.append(page_number)
        if current is not None:
            notices.append(_interpret_notice(current, pages, len(notices) + 1))
    if not notices:
        raise ValueError("No reconstructable change notices")
    first = page_texts[1]
    coverage = re.search(r"Avser perioden:\s*([^\n]+)", first)
    if not coverage:
        raise ValueError("Missing report coverage period")
    heading_position = first.find("Myndighet Produktkod")
    preamble = first[:heading_position]
    document_date = next(
        (
            line
            for line in preamble.splitlines()
            if re.fullmatch(r"[A-Za-zåäöÅÄÖ]+ \d{4}", line)
        ),
        None,
    )
    return ChangesReport(
        provenance=snapshot.provenance(
            canonical=snapshot.final_url,
            selector="PDF words grouped by printed columns",
        ),
        link_label=link_label or coverage[1],
        title="Förändringar i den officiella statistiken",
        document_date_text=document_date,
        coverage_period_text=coverage[1],
        page_count=len(page_texts),
        document_notes=[preamble],
        partial_sample=False,
        notices=notices,
        page_texts=page_texts,
    )


@dataclass(frozen=True)
class DownloadsAdapter:
    source: str

    def collect(self, context: ExtractionContext) -> SourceResult:
        result = SourceResult(source=self.source, discovered=["index"])
        try:
            snapshot = context.fetch(FetchRequest(url=URLS[self.source]))
            if self.source == "official_products":
                soup = _html(snapshot)
                anchors = [
                    a
                    for a in soup.select("main a[href]")
                    if _text(a) == "Statistikprodukter inom den officiella statistiken"
                ]
                if len(anchors) != 1 or not urlsplit(
                    anchors[0]["href"]
                ).path.lower().endswith(".xlsx"):
                    raise ValueError(
                        "Current official-product workbook anchor missing or ambiguous"
                    )
                workbook = context.fetch(
                    FetchRequest(
                        url=absolute_url(snapshot.final_url, anchors[0]["href"])
                    )
                )
                result.documents["index"] = parse_workbook(workbook)
            elif self.source == "hvd":
                result.documents["index"] = parse_hvd(snapshot)
            elif self.source == "economy":
                collection = parse_economy(snapshot)
                result.documents["index"] = collection
                entries = []
                for entry in collection.entries:
                    key = url_key(entry.url)
                    result.discovered.append(key)
                    try:
                        detail = context.fetch(FetchRequest(url=entry.url))
                        result.documents[key] = parse_economy_detail(detail, entry)
                        entries.append(result.documents[key])
                    except SOURCE_ERRORS as exc:
                        result.failures[key] = str(exc)
                        entries.append(entry)
                result.documents["index"] = collection.model_copy(
                    update={"entries": entries}
                )
            else:
                inventory = parse_changes_index(snapshot)
                result.documents["index"] = inventory
                for entry in inventory.entries:
                    result.discovered.append(entry.key)
                    try:
                        report = context.fetch(FetchRequest(url=entry.url))
                        result.documents[entry.key] = parse_changes(
                            report, link_label=entry.title
                        )
                    except SOURCE_ERRORS as exc:
                        result.failures[entry.key] = str(exc)
                if result.failures:
                    result.documents["index"] = inventory.model_copy(
                        update={"complete": False}
                    )
        except SOURCE_ERRORS as exc:
            result.failures["index"] = str(exc)
            result.discovery_complete = False
        return result


ADAPTERS = {name: DownloadsAdapter(name) for name in URLS}
