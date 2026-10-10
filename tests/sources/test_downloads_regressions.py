"""Offline regressions against archived sources and explicit failure simulations."""

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scb_extract.core import ExtractionContext, url_key
from scb_extract.sources.downloads import (
    ADAPTERS,
    parse_changes,
    parse_workbook,
)


def test_all_download_adapters_replay_with_complete_inventory():
    with ExtractionContext(
        Path(__file__).parents[1] / "fixtures/raw", offline=True
    ) as context:
        for source, expected in [
            ("official_products", 1),
            ("hvd", 1),
            ("changes", 5),
        ]:
            result = ADAPTERS[source].collect(context)
            assert result.discovery_complete
            assert not result.failures
            assert len(result.documents) == len(result.discovered) == expected
        hvd = ADAPTERS["hvd"].collect(context).documents["index"]
        assert hvd.product_links_only
        assert all(link.product_code for group in hvd.groups for link in group.links)


def test_workbook_preserves_whitespace_and_rejects_changed_headers(snapshot):
    original = snapshot("official_products_0")
    book = load_workbook(io.BytesIO(original.content))
    sheet = book.active
    sheet.cell(2, 3).value = "A title with trailing whitespace  "
    buffer = io.BytesIO()
    book.save(buffer)
    parsed = parse_workbook(original.model_copy(update={"content": buffer.getvalue()}))
    assert parsed.rows[0].product_name == "A title with trailing whitespace  "
    sheet.cell(1, 5).value = "Användare"
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    with pytest.raises(ValueError, match="headers changed"):
        parse_workbook(original.model_copy(update={"content": buffer.getvalue()}))


def test_pdf_all_rows_cross_page_prose_and_repeated_codes(snapshot):
    for i, count in enumerate([10, 18, 24, 14]):
        report = parse_changes(snapshot(f"changes_{i}"))
        assert len(report.notices) == count
        assert set(report.page_texts) == set(range(1, report.page_count + 1))
        assert [n.row_ordinal for n in report.notices] == list(range(1, count + 1))
    report = parse_changes(snapshot("changes_1"))
    repeated = [n for n in report.notices if n.product_code == "SF0206"]
    assert len(repeated) == 2
    assert repeated[0].pages == [1, 2]
    assert "kvaliteteten inte var tillräckligt bra" in repeated[0].description
    assert repeated[0].timing_claim == "planned"
    assert repeated[0].effective_date is None
    merge = next(
        n
        for n in parse_changes(snapshot("changes_3")).notices
        if n.product_code == "NV0124"
    )
    assert merge.pages == [2, 3]
    assert merge.related_product_codes == ["NV0801", "NV0123"]
    assert merge.effective_date is None
    assert merge.review_required
    assert "referen- speriod kvartal 1 2024" in merge.description


def test_inactive_closed_and_merge_are_distinct_source_claims(snapshot):
    inactive = next(
        n
        for n in parse_changes(snapshot("changes_0")).notices
        if n.product_code == "AM0206"
    )
    assert inactive.change_kinds == ["inactive"]
    notices = parse_changes(snapshot("changes_1")).notices
    assert next(n for n in notices if n.product_code == "MI0308").change_kinds == [
        "closed"
    ]
    assert next(n for n in notices if n.product_code == "EN0113").change_kinds == [
        "merge"
    ]


def test_child_failure_is_keyed_and_inventory_marked_incomplete(snapshot):
    index_snapshot = snapshot("changes")

    class FailingChildren:
        def fetch(self, request):
            if request.url == index_snapshot.request.url:
                return index_snapshot
            raise OSError("Simulated unavailable report")

    result = ADAPTERS["changes"].collect(FailingChildren())
    assert result.discovery_complete
    assert result.documents["index"].complete is False
    assert len(result.failures) == 4
    assert all(
        entry.key == url_key(entry.url) for entry in result.documents["index"].entries
    )


def test_current_workbook_discovered_from_label_not_pinned_filename(snapshot):
    page = snapshot("official_products")
    original = snapshot("official_products_0")
    old_path = original.request.url.removeprefix("https://www.scb.se")
    new_path = "/globalassets/new-current-workbook.xlsx"
    page = page.model_copy(
        update={"content": page.content.replace(old_path.encode(), new_path.encode())}
    )
    requests = []

    class CurrentWorkbook:
        def fetch(self, request):
            requests.append(request.url)
            return page if request.url == page.request.url else original

    result = ADAPTERS["official_products"].collect(CurrentWorkbook())
    assert not result.failures
    assert requests == [page.request.url, "https://www.scb.se" + new_path]


def test_corrupt_pdf_response_fails_each_discovered_document(snapshot):
    page = snapshot("changes")
    corrupt = snapshot("changes_0").model_copy(update={"content": b"not a PDF"})

    class CorruptReports:
        def fetch(self, request):
            return page if request.url == page.request.url else corrupt

    result = ADAPTERS["changes"].collect(CorruptReports())
    assert len(result.failures) == 4
    assert list(result.documents) == ["index"]
    assert result.documents["index"].complete is False
